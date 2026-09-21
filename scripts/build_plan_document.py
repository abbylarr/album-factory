from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[1]
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin = section.bottom_margin = Cm(1.8)
section.left_margin = section.right_margin = Cm(2.1)
normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.12
for name, size in [("Title", 28), ("Subtitle", 16), ("Heading 1", 17), ("Heading 2", 12)]:
    style = doc.styles[name]
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(14)
    style.paragraph_format.space_after = Pt(7)
    style.paragraph_format.keep_with_next = True
for line in (root / "docs/product-plan.md").read_text().splitlines():
    if not line:
        continue
    if line.startswith("# "):
        doc.add_paragraph(line[2:], "Title")
    elif line == "## План продукта и реализации":
        doc.add_paragraph(line[3:], "Subtitle")
    elif line.startswith("## "):
        doc.add_paragraph(line[3:], "Heading 1")
    elif line.startswith("### "):
        doc.add_paragraph(line[4:], "Heading 2")
    else:
        p = doc.add_paragraph(line)
        p.paragraph_format.widow_control = True
footer = section.footer.paragraphs[0]
footer.alignment = 2
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)
doc.core_properties.title = "Album Factory План продукта и реализации"
doc.core_properties.author = ""
# The runtime's default Word template carries a title border in its styles.
for container in (doc.styles.element, doc.element):
    for border in list(container.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
output = root / "output/documents"
output.mkdir(parents=True, exist_ok=True)
doc.save(output / "Album-Factory-Plan.docx")
print(output / "Album-Factory-Plan.docx")
